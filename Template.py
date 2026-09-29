
import os
import re
from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

# Cột chứa mã CC : cột chứa Repeat tương ứng
CC_COLUMNS = {
    "AA": "AB",
    "AD": "AE",
    "AG": "AH",
    "AJ": "AK",
}

START_ROW = 3
END_ROW = 102

# Các ký tự được phép sử dụng trong mã CC
VALID_CHARS = set(
    "0123456789ABCDEFGHJKLMNPQRSTUVWXYZ"
)

EXPECTED_CC_COUNT = 400


# ============================================================
# 1. LOAD EXCEL TEMPLATE
# ============================================================

def load_excel_template(excel_path, sheet_name=None):
    """
    Mở file Excel template .xlsm.

    Parameters
    ----------
    excel_path : str
        Đường dẫn đến file .xlsm.

    sheet_name : str or None
        Tên sheet cần sử dụng.
        Nếu None -> sử dụng sheet active.

    Returns
    -------
    wb : openpyxl Workbook
    ws : openpyxl Worksheet
    """

    # --------------------------------------------------------
    # Kiểm tra file tồn tại
    # --------------------------------------------------------

    if not os.path.isfile(excel_path):
        raise FileNotFoundError(
            f"Không tìm thấy file Excel:\n{excel_path}"
        )

    # --------------------------------------------------------
    # Kiểm tra extension
    # --------------------------------------------------------

    extension = os.path.splitext(excel_path)[1].lower()

    if extension != ".xlsm":
        raise ValueError(
            f"File template phải có định dạng .xlsm\n"
            f"File hiện tại: {extension}"
        )

    # --------------------------------------------------------
    # Mở workbook
    # --------------------------------------------------------

    wb = load_workbook(
        excel_path,
        keep_vba=True,
        data_only=False
    )

    # --------------------------------------------------------
    # Chọn worksheet
    # --------------------------------------------------------

    if sheet_name is None:
        ws = wb.active

    else:
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"Không tìm thấy sheet '{sheet_name}'.\n"
                f"Các sheet hiện có: {wb.sheetnames}"
            )

        ws = wb[sheet_name]

    return wb, ws


# ============================================================
# 2. READ CC CODES
# ============================================================

def read_cc_codes(ws):
    """
    Đọc toàn bộ mã CC từ 4 vùng:

        AA3:AA102 -> Repeat AB
        AD3:AD102 -> Repeat AE
        AG3:AG102 -> Repeat AH
        AJ3:AJ102 -> Repeat AK

    Returns
    -------
    records : list of dict
    """

    records = []

    for code_col, repeat_col in CC_COLUMNS.items():

        for row in range(START_ROW, END_ROW + 1):

            code_cell = f"{code_col}{row}"
            repeat_cell = f"{repeat_col}{row}"

            value = ws[code_cell].value

            # ------------------------------------------------
            # Bỏ qua ô trống
            # ------------------------------------------------

            if value is None:
                continue

            # ------------------------------------------------
            # Chuẩn hóa dữ liệu
            # ------------------------------------------------

            code = str(value).strip().upper()

            records.append({
                "code": code,
                "code_cell": code_cell,
                "repeat_cell": repeat_cell,
                "code_column": code_col,
                "repeat_column": repeat_col,
                "row": row,
            })

    return records


# ============================================================
# 3. VALIDATE CC CODES
# ============================================================

def validate_cc_codes(records):
    """
    Kiểm tra danh sách mã CC trong template.

    Kiểm tra:
        1. Có đúng 400 mã hay không
        2. Mỗi mã có đúng 2 ký tự
        3. Ký tự thuộc tập cho phép
        4. Không có mã trùng
    """

    errors = []

    # --------------------------------------------------------
    # Kiểm tra số lượng
    # --------------------------------------------------------

    if len(records) != EXPECTED_CC_COUNT:
        errors.append(
            f"Số lượng mã CC = {len(records)}, "
            f"nhưng yêu cầu = {EXPECTED_CC_COUNT}."
        )

    # --------------------------------------------------------
    # Kiểm tra từng mã
    # --------------------------------------------------------

    seen_codes = {}

    for record in records:

        code = record["code"]
        cell = record["code_cell"]

        # ----------------------------------------------------
        # Độ dài
        # ----------------------------------------------------

        if len(code) != 2:

            errors.append(
                f"{cell}: mã '{code}' không có đúng 2 ký tự."
            )

            continue

        # ----------------------------------------------------
        # Kiểm tra ký tự
        # ----------------------------------------------------

        invalid_chars = [
            char
            for char in code
            if char not in VALID_CHARS
        ]

        if invalid_chars:

            errors.append(
                f"{cell}: mã '{code}' chứa ký tự "
                f"không hợp lệ: {invalid_chars}"
            )

        # ----------------------------------------------------
        # Kiểm tra mã trùng
        # ----------------------------------------------------

        if code in seen_codes:

            previous_cell = seen_codes[code]

            errors.append(
                f"Mã '{code}' bị trùng tại "
                f"{previous_cell} và {cell}."
            )

        else:
            seen_codes[code] = cell

    # --------------------------------------------------------
    # Nếu có lỗi -> dừng chương trình
    # --------------------------------------------------------

    if errors:

        error_text = "\n".join(
            f"  - {error}"
            for error in errors
        )

        raise ValueError(
            "Template CC không hợp lệ:\n"
            + error_text
        )

    return True


# ============================================================
# 4. BUILD CC LOOKUP
# ============================================================

def build_cc_lookup(records):
    """
    Tạo dictionary để tra cứu nhanh:

        CC code -> vị trí ô Code + ô Repeat

    Ví dụ:

        "74" -> {
            "code_cell": "AD5",
            "repeat_cell": "AE5"
        }
    """

    cc_lookup = {}

    for record in records:

        code = record["code"]

        cc_lookup[code] = {
            "code_cell": record["code_cell"],
            "repeat_cell": record["repeat_cell"],
            "code_column": record["code_column"],
            "repeat_column": record["repeat_column"],
            "row": record["row"],
        }

    return cc_lookup


# ============================================================
# 5. PRINT TEMPLATE SUMMARY
# ============================================================

def print_template_summary(records, cc_lookup):
    """
    In thông tin kiểm tra template ra console.
    """

    print("=" * 60)
    print("CC TEMPLATE SUMMARY")
    print("=" * 60)

    print(f"Total CC codes : {len(records)}")
    print(f"Lookup size    : {len(cc_lookup)}")

    print("\nCC distribution:")

    for code_col, repeat_col in CC_COLUMNS.items():

        count = sum(
            1
            for record in records
            if record["code_column"] == code_col
        )

        print(
            f"  {code_col} -> {repeat_col} : "
            f"{count} codes"
        )

    print("=" * 60)


# ============================================================
# 6. TEST FUNCTION - STAGE 1
# ============================================================

def test_stage_1(excel_path, sheet_name=None):

    # --------------------------------------------------------
    # Load Excel
    # --------------------------------------------------------

    wb, ws = load_excel_template(
        excel_path,
        sheet_name=sheet_name
    )

    print(f"\nExcel file : {excel_path}")
    print(f"Worksheet  : {ws.title}")

    # --------------------------------------------------------
    # Read CC
    # --------------------------------------------------------

    records = read_cc_codes(ws)

    # --------------------------------------------------------
    # Validate
    # --------------------------------------------------------

    validate_cc_codes(records)

    # --------------------------------------------------------
    # Build lookup
    # --------------------------------------------------------

    cc_lookup = build_cc_lookup(records)

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    print_template_summary(
        records,
        cc_lookup
    )

    # --------------------------------------------------------
    # In thử một số mã
    # --------------------------------------------------------

    print("\nFirst 10 CC codes:")

    for record in records[:10]:

        print(
            f"  {record['code']}  | "
            f"{record['code_cell']} -> "
            f"{record['repeat_cell']}"
        )

    print("\nStage 1 completed successfully.")

    return wb, ws, records, cc_lookup

if __name__ == "__main__":

    excel_path = r"E:\Your_Project\CC_Template.xlsm"

    wb, ws, records, cc_lookup = test_stage_1(
        excel_path
    )
